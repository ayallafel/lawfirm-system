import streamlit as st
import pandas as pd
from docx import Document
import os
import io
import datetime
import time

# --- הגדרות עיצוב וממשק ---
st.set_page_config(page_title="אילן פלדמן - משרד עורכי דין", page_icon="🏢", layout="centered")

st.markdown("""
    <style>
    .stApp {
        background-color: #fafafa;
        background-image: 
            linear-gradient(90deg, rgba(0,0,0,0.03) 1px, transparent 1px),
            linear-gradient(0deg, rgba(0,0,0,0.03) 1px, transparent 1px);
        background-size: 40px 40px;
    }
    .main .block-container, p, span, label, div {
        direction: rtl;
        text-align: right;
    }
    .stButton button {
        width: 100%;
        border-radius: 8px;
        font-weight: bold;
    }
    [data-testid="stSidebar"] {
        direction: rtl;
    }
    [data-testid="stSidebar"] * {
        text-align: right;
    }
    </style>
""", unsafe_allow_html=True)

LOGO_FILE = "logo.png"

# --- ניהול מצב האפליקציה (Session State) לבחירת פרויקט ---
if 'selected_project' not in st.session_state:
    st.session_state.selected_project = None

PROJECTS = {
    "סמילנסקי": "smilanski",
    "תרנא": "tarna",
    "אקווה": "aqua",
    "ביאליק": "bialik"
}

# ==================== מסך הבית: בחירת פרויקט ====================
if st.session_state.selected_project is None:
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        if os.path.exists(LOGO_FILE):
            st.image(LOGO_FILE, use_container_width=True)
        else:
            st.markdown("<h2 style='text-align: center;'>אילן פלדמן - משרד עורכי דין</h2>", unsafe_allow_html=True)

    st.markdown("<h1 style='text-align: center;'>בחר פרויקט לניהול</h1>", unsafe_allow_html=True)
    st.markdown("<p style='text-align: center; color: #555;'>לחץ על אחד הפרויקטים כדי להיכנס למערכת הניהול והפקת החוזים שלו</p>", unsafe_allow_html=True)
    st.markdown("---")

    for heb_name, code in PROJECTS.items():
        col_img, col_btn = st.columns([1, 3])
        
        with col_img:
            building_img_path = os.path.join("projects", code, "building.png")
            if os.path.exists(building_img_path):
                st.image(building_img_path, width=100)
            else:
                st.markdown("🏢 **[תמונת בניין]**")
                
        with col_btn:
            st.markdown(f"### פרויקט {heb_name}")
            if st.button(f"כניסה לפרויקט {heb_name}", key=f"btn_{code}"):
                st.session_state.selected_project = (heb_name, code)
                st.rerun()
        st.markdown("---")

# ==================== מסך פרויקט פעיל ====================
else:
    selected_project_heb, project_code = st.session_state.selected_project

    if st.button("⬅️ חזרה לבחירת פרויקטים"):
        st.session_state.selected_project = None
        st.rerun()

    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        if os.path.exists(LOGO_FILE):
            st.image(LOGO_FILE, use_container_width=True)

    st.markdown(f"<h1 style='text-align: center;'>פרויקט: {selected_project_heb}</h1>", unsafe_allow_html=True)
    st.markdown("---")

    # הגדרת תיקיות ונתיבים דינאמיים לפרויקט הנבחר
    PROJECT_DIR = os.path.join("projects", project_code)
    os.makedirs(PROJECT_DIR, exist_ok=True)
    os.makedirs(os.path.join(PROJECT_DIR, "templates"), exist_ok=True)

    DB_FILE = os.path.join(PROJECT_DIR, "clients.xlsx")
    TEMPLATE_FILE = os.path.join(PROJECT_DIR, "templates", "lease_template.docx")
    REG_TEMPLATE_FILE = os.path.join(PROJECT_DIR, "templates", "registration_form.docx")

    if not os.path.exists(TEMPLATE_FILE):
        doc = Document()
        doc.add_paragraph("בניין: {{מספר_בניין}}\nדירה: {{מספר_דירה}}\nקונה 1: {{שם_הקונה_1}}\nת.ז 1: {{תז_קונה_1}}\nמחיר: {{מחיר_הדירה}}")
        doc.save(TEMPLATE_FILE)

    if not os.path.exists(REG_TEMPLATE_FILE):
        doc = Document()
        doc.add_paragraph("מספר בניין:\nמספר דירה:\nקומה:\nמספר חדרים:\nשם הקונה 1:\nת\"ז קונה 1:\nשם הקונה 2:\nת\"ז קונה 2:\nכתובת הקונה:\nמספר טלפון:\nדוא\"ל:\nמחיר הדירה:\nמחיר הדירה במילים:")
        doc.save(REG_TEMPLATE_FILE)

    def get_current_template_version():
        if os.path.exists(TEMPLATE_FILE):
            mtime = os.path.getmtime(TEMPLATE_FILE)
            return time.strftime("%Y-%m-%d %H:%M", time.localtime(mtime))
        return datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

    def extract_data_smart(doc):
        """חילוץ דינאמי וחכם של כל שדה מתוך טופס הוורד"""
        data = {}
        for para in doc.paragraphs:
            text = para.text.strip()
            if ":" in text:
                parts = text.split(":", 1)
                key = parts[0].replace("שדה", "").strip()
                val = parts[1].strip()
                if key and val:
                    data[key] = val
        return data

    def generate_contract(data):
        doc = Document(TEMPLATE_FILE)
        
        # מיפוי חכם של שדות מתוך הנתונים אל המשתנים החדשים בחוזה
        now = datetime.datetime.now()
        placeholders = {
            "{{מספר_בניין}}": str(data.get("מספר בניין", data.get("בניין", "1"))),
            "{{מספר_דירה}}": str(data.get("מספר דירה", data.get("דירה", "1"))),
            "{{קומה}}": str(data.get("קומה", "1")),
            "{{מספר_חדרים}}": str(data.get("מספר חדרים", data.get("חדרים", "4"))),
            "{{יום_חתימה}}": str(data.get("יום חתימה", now.strftime("%d"))),
            "{{חודש_חתימה}}": str(data.get("חודש חתימה", now.strftime("%m"))),
            "{{שנת_חתימה}}": str(data.get("שנת חתימה", now.strftime("%Y"))),
            "{{שם_הקונה_1}}": str(data.get("שם הקונה 1", data.get("שם הקונה", ""))),
            "{{תז_קונה_1}}": str(data.get("ת\"ז קונה 1", data.get("ת\"ז קונה", ""))),
            "{{שם_הקונה_2}}": str(data.get("שם הקונה 2", "")),
            "{{תז_קונה_2}}": str(data.get("ת\"ז קונה 2", "")),
            "{{כתובת_הקונה}}": str(data.get("כתובת הקונה", data.get("כתובת נוכחית", ""))),
            "{{טלפון_הקונה_1}}": str(data.get("מספר טלפון", data.get("טלפון", ""))),
            "{{דואל_הקונה_1}}": str(data.get("דוא\"ל", data.get("אימייל", ""))),
            "{{מחיר_הדירה}}": str(data.get("מחיר הדירה", data.get("מחיר הנכס", ""))),
            "{{מחיר_הדירה_במילים}}": str(data.get("מחיר הדירה במילים", "")),
        }

        # החלפה בכל פסקאות המסמך
        for para in doc.paragraphs:
            for key, val in placeholders.items():
                if key in para.text:
                    para.text = para.text.replace(key, val)
                    
        # בדיקה גם בטבלאות אם קיימות במסמך
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for key, val in placeholders.items():
                        if key in cell.text:
                            cell.text = cell.text.replace(key, val)

        bio = io.BytesIO()
        doc.save(bio)
        return bio.getvalue()

    def process_and_download(extracted_data):
        current_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        template_ver = get_current_template_version()
        
        extracted_data["תאריך הוספה"] = current_time
        extracted_data["גרסת חוזה"] = template_ver
        
        new_row = pd.DataFrame([extracted_data])
        id_field = 'ת"ז קונה 1' if 'ת"ז קונה 1' in extracted_data else ('ת"ז קונה' if 'ת"ז קונה' in extracted_data else list(extracted_data.keys())[1])

        if os.path.exists(DB_FILE):
            df_existing = pd.read_excel(DB_FILE)
            if id_field in df_existing.columns and str(extracted_data.get(id_field)) in df_existing[id_field].astype(str).values:
                for col in extracted_data.keys():
                    if col in df_existing.columns:
                        df_existing.loc[df_existing[id_field].astype(str) == str(extracted_data.get(id_field)), col] = extracted_data[col]
                df_existing.to_excel(DB_FILE, index=False)
                st.success(f"הלקוח כבר קיים במאגר – פרטיו עודכנו בהצלחה לגרסה העדכנית ({template_ver})!")
            else:
                df_updated = pd.concat([df_existing, new_row], ignore_index=True, sort=False)
                df_updated.to_excel(DB_FILE, index=False)
                st.success(f"הקונה החדש נוסף בהצלחה למאגר של {selected_project_heb}.")
        else:
            new_row.to_excel(DB_FILE, index=False)
            st.success(f"נוצר מאגר נתונים חדש עבור פרויקט {selected_project_heb}.")
        
        contract_bytes = generate_contract(extracted_data)
        buyer_name = str(extracted_data.get("שם הקונה 1", extracted_data.get("שם הקונה", "לקוח_ללא_שם"))).replace(" ", "_")
        filename = f"הסכם_מכר_{project_code}_{buyer_name}.docx"
        
        st.download_button(
            label="📥 הורד את הסכם המכר המוכן למחשב",
            data=contract_bytes,
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    # --- תפריט צדדי (הגדרות מאסטר לפרויקט הנוכחי) ---
    with st.sidebar:
        st.header(f"⚙️ הגדרות: {selected_project_heb}")
        st.markdown("ניהול תבניות המאסטר של פרויקט זה.")
        
        st.markdown("---")
        st.subheader("טופס הרשמה")
        if os.path.exists(REG_TEMPLATE_FILE):
            with open(REG_TEMPLATE_FILE, "rb") as reg_f:
                st.download_button(
                    label="📥 הורד טופס הרשמה ריק",
                    data=reg_f,
                    file_name=f"registration_form_{project_code}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
        uploaded_reg_template = st.file_uploader("החלף טופס הרשמה מאסטר", type=["docx"], key="master_reg")
        if uploaded_reg_template is not None:
            if st.button("שמור טופס הרשמה חדש"):
                try:
                    with open(REG_TEMPLATE_FILE, "wb") as f:
                        f.write(uploaded_reg_template.getbuffer())
                    st.success("טופס ההרשמה עודכן!")
                except Exception as e:
                    st.error(f"שגיאה: {e}")

        st.markdown("---")
        st.subheader("חוזה מכר (מאסטר)")
        st.caption(f"גרסת מאסטר נוכחית: {get_current_template_version()}")
        
        uploaded_template = st.file_uploader("העלה חוזה מאסטר חדש", type=["docx"], key="master_lease")
        
        if uploaded_template is not None:
            if st.button("שמור תבנית חוזה מאסטר חדשה"):
                try:
                    with open(TEMPLATE_FILE, "wb") as f:
                        f.write(uploaded_template.getbuffer())
                    st.success("תבנית החוזה עודכנה בהצלחה!")
                    time.sleep(1)
                    st.rerun()
                except Exception as e:
                    st.error(f"שגיאה: {e}")

    # --- ממשק ראשי לפרויקט ---
    st.subheader("📌 הוספת לקוח חדש")
    method = st.radio("בחר את אופן הוספת הלקוח:", ["העלאת טופס הרשמה", "הזנה ידנית"], horizontal=True)

    if method == "העלאת טופס הרשמה":
        uploaded_file = st.file_uploader("העלי את טופס ההרשמה של הקונה (קובץ Word)", type=["docx"], key="reg_form")
        
        if uploaded_file is not None:
            if st.button("חלץ נתונים והפק הסכם מכר"):
                try:
                    doc = Document(uploaded_file)
                    data = extract_data_smart(doc)
                    if not data:
                        st.error("לא הצלחתי למצוא נתונים בפורמט מפתח: ערך.")
                    else:
                        process_and_download(data)
                except Exception as e:
                    st.error(f"אירעה שגיאה במהלך העיבוד: {e}")

    else:
        with st.form("manual_entry_form"):
            col_a, col_b = st.columns(2)
            with col_a:
                building_num = st.text_input("מספר בניין", value="1")
                apartment_num = st.text_input("מספר דירה")
                floor = st.text_input("קומה")
                rooms = st.text_input("מספר חדרים", value="4")
            with col_b:
                buyer1_name = st.text_input("שם הקונה 1")
                buyer1_id = st.text_input('ת"ז קונה 1')
                buyer_phone = st.text_input("מספר טלפון")
                buyer_email = st.text_input("דוא\"ל")
            
            price = st.text_input("מחיר הדירה (₪)")
            price_words = st.text_input("מחיר הדירה במילים (לדוגמה: שני מיליון ומאה אלף)")
            
            submitted = st.form_submit_button("שמור נתונים והפק הסכם מכר")
            
            if submitted:
                if not buyer1_name or not buyer1_id or not apartment_num:
                    st.error("חובה להזין לפחות מספר דירה, שם קונה ותעודת זהות.")
                else:
                    manual_data = {
                        "מספר בניין": building_num,
                        "מספר דירה": apartment_num,
                        "קומה": floor,
                        "מספר חדרים": rooms,
                        "שם הקונה 1": buyer1_name,
                        "ת\"ז קונה 1": buyer1_id,
                        "מספר טלפון": buyer_phone,
                        "דוא\"ל": buyer_email,
                        "מחיר הדירה": price,
                        "מחיר הדירה במילים": price_words
                    }
                    process_and_download(manual_data)

    # --- רשימת הלקוחות בפרויקט כטבלה נקייה מסודרת ---
    st.markdown("---")
    st.markdown(f"<h2>📋 רשימת הלקוחות בפרויקט: {selected_project_heb}</h2>", unsafe_allow_html=True)

    if os.path.exists(DB_FILE):
        df = pd.read_excel(DB_FILE)
        
        if "תאריך הוספה" not in df.columns:
            df["תאריך הוספה"] = "לא ידוע"
        if "גרסת חוזה" not in df.columns:
            df["גרסת חוזה"] = "ישנה"

        st.dataframe(df, use_container_width=True)

        st.markdown("### 📥 הורדת חוזים ועדכון גרסאות ללקוחות")
        
        for idx, row in df.iterrows():
            client_data = row.to_dict()
            buyer_name_str = str(client_data.get('שם הקונה 1', client_data.get('שם הקונה', 'לקוח')))
            buyer_id_str = str(client_data.get('ת"ז קונה 1', client_data.get('ת"ז קונה', '')))
            
            cols = st.columns([3, 1, 1])
            with cols[0]:
                st.text(f"{buyer_name_str} (ת\"ז: {buyer_id_str}) | גרסה: {client_data.get('גרסת חוזה', '')}")
            
            with cols[1]:
                contract_bytes = generate_contract(client_data)
                filename = f"הסכם_מכר_{project_code}_{buyer_name_str.replace(' ', '_')}.docx"
                st.download_button(
                    label="📥 הורד חוזה",
                    data=contract_bytes,
                    file_name=filename,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key=f"download_{project_code}_{idx}"
                )
                
            with cols[2]:
                if st.button("🔄 עדכן לגרסה חדשה", key=f"update_ver_{project_code}_{idx}"):
                    current_ver = get_current_template_version()
                    df.loc[df.iloc[:, 1].astype(str) == buyer_id_str, 'גרסת חוזה'] = current_ver
                    df.to_excel(DB_FILE, index=False)
                    st.success(f"החוזה עודכן לגרסה {current_ver}!")
                    st.rerun()

        st.markdown("---")
        with open(DB_FILE, "rb") as f:
            st.download_button(
                label=f"📊 הורד את טבלת האקסל המלאה של {selected_project_heb} למחשב",
                data=f,
                file_name=f"clients_{project_code}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"excel_full_{project_code}"
            )
    else:
        st.info(f"עדיין אין קונים במאגר של פרויקט {selected_project_heb}. הטבלה תיווצר אוטומטית כשתזיני את הלקוח הראשון.")